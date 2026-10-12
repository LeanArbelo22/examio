from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path('docs/documentos/Arquitectura_del_Sistema_Examio.docx')
doc = Document()
s = doc.sections[0]
s.page_width = Cm(21); s.page_height = Cm(29.7)
s.top_margin = Cm(2); s.bottom_margin = Cm(2)
s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Number']:
    st = doc.styles[name]
    st.font.name = 'Calibri'; st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.08
    if st.element.rPr is not None:
        for color in st.element.rPr.findall(qn('w:color')):
            for attr in ['themeColor', 'themeTint', 'themeShade']:
                color.attrib.pop(qn('w:'+attr), None)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'es-AR'); doc.styles['Normal'].element.get_or_add_rPr().append(lang)
doc.core_properties.title = 'Arquitectura del Sistema Examio'
doc.core_properties.subject = 'Componentes y funcionamiento del sistema'
doc.core_properties.author = ''
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def p(t, style=None): return doc.add_paragraph(t, style)
def h(t): return doc.add_heading(t, 1)
def h2(t): return doc.add_heading(t, 2)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for cell,txt in zip(t.rows[0].cells,headers): cell.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); mar.append(el)
            pr.append(mar)
            if i==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EDEDED'); pr.append(sh)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2); par.paragraph_format.space_before=Pt(2)
                par.paragraph_format.line_spacing=1.0
                for run in par.runs: run.font.size=Pt(10.5); run.bold=i==0
    return t

p('Arquitectura del Sistema Examio', 'Title')
p('Una explicación de sus partes y de cómo trabajan juntas', 'Subtitle')
p('7 de octubre de 2026')
p('Examio es un sistema para administrar evaluaciones académicas. Permite organizar materias y alumnos, preparar exámenes, rendirlos, corregir respuestas y consultar resultados. Este documento explica cómo se conectan sus componentes y qué responsabilidad tiene cada uno.')
h('Cómo se organiza el sistema')
p('La aplicación web se apoya en una arquitectura por capas. La interfaz muestra la información y recibe las acciones del usuario; la API procesa las operaciones y controla los permisos; SQL Server guarda los datos y aplica restricciones de integridad. Next.js también aporta una parte de servidor que administra la sesión y sirve de puente hacia la API.')
flow=p('Usuario en el navegador\n↕\nAplicación web y servidor Next.js\n↕\nAPI de Examio en Node.js y Express\n↕\nBase de datos SQL Server')
flow.alignment=WD_ALIGN_PARAGRAPH.CENTER
flow.paragraph_format.space_before=Pt(4); flow.paragraph_format.space_after=Pt(10)
for run in flow.runs: run.bold=True
p('Las flechas representan solicitudes y respuestas. El navegador no se conecta directamente a la base de datos. Dentro del servidor Next.js, el puente hacia la API se llama BFF, abreviatura de Backend for Frontend.')
h('Quiénes lo utilizan')
table(['Rol','Uso principal'],[
('Administrador','Gestiona usuarios, alumnos, docentes, carreras, materias y comisiones.'),
('Docente','Prepara y publica evaluaciones, reutiliza preguntas, corrige y consulta reportes.'),
('Alumno','Accede a sus evaluaciones, responde preguntas y consulta resultados.')
],[3.3,13.3])

doc.add_page_break()
h('Las partes de la arquitectura')
h2('La aplicación web')
p('El frontend está construido con Next.js 16, React 19 y TypeScript. Next.js organiza las páginas y la navegación; React permite componer pantallas con componentes reutilizables; TypeScript ayuda a detectar incompatibilidades en los datos durante el desarrollo. Tailwind CSS se usa para los estilos.')
p('Las rutas principales se separan por rol: /admin, /profesor y /alumno. Cada área reúne las funciones que necesita ese usuario. Los formularios validan la entrada para dar mensajes claros, pero la API vuelve a validar antes de aceptar una operación.')
p('El código también se divide según su tarea. Las páginas están en src/app, los componentes en src/components, las consultas y actualizaciones en src/hooks, las llamadas HTTP en src/services/api y los contratos de datos en src/types. Las herramientas comunes, la sesión y el cliente HTTP están en src/lib.')
h2('El puente entre la web y la API')
p('Las consultas de la interfaz pasan por hooks de TanStack React Query, servicios de dominio y un cliente común llamado apiFetch. Los hooks son funciones que conectan las pantallas con sus datos. React Query administra la carga, los errores y la caché, que permite reutilizar datos ya consultados.')
p('apiFetch envía las solicitudes a /api/examio/*, dentro del mismo sitio. El BFF toma las credenciales de la sesión en el servidor y agrega la autorización al comunicarse con la API. Así, los componentes del navegador no necesitan manejar los tokens de acceso.')
h2('La API y la base de datos')
p('La API REST está desarrollada con Node.js, Express 5 y TypeScript. Recibe solicitudes HTTP con datos en JSON, valida su contenido con Zod y comprueba el rol y los permisos del usuario. Sus rutas y controladores resuelven operaciones como publicar una evaluación, guardar una respuesta o consultar una nota.')
p('La conexión con SQL Server se realiza mediante mssql, que reutiliza conexiones disponibles. La base GestionExamio guarda la información académica, las relaciones entre registros y parte de la lógica mediante procedimientos almacenados. Las operaciones que deben completarse juntas utilizan transacciones: si una parte falla, se revierten los cambios de esa operación.')
h2('Tecnologías de apoyo')
p('NextAuth administra la sesión web. Zod valida formularios y contratos según la capa. Recharts dibuja los gráficos de los reportes. IndexedDB, un almacenamiento local del navegador, conserva respuestas de examen pendientes de sincronización. Cada herramienta cubre una necesidad concreta dentro del sistema.')

doc.add_page_break()
h('Cómo se relacionan los datos')
p('MateriaDivision es una entidad central: vincula una materia con una división o comisión y, cuando corresponde, con un docente. Las inscripciones determinan qué alumnos pertenecen a esa comisión y qué evaluaciones pueden consultar.')
table(['Concepto','Qué representa'],[
('Actividad','El contenido de una evaluación, trabajo práctico o tarea, con sus grupos de preguntas.'),
('Publicación','La disponibilidad de una actividad para una comisión, con fechas y visibilidad.'),
('Intento','La realización de una publicación por un alumno en particular.'),
('Pregunta y opción','Los enunciados y las alternativas disponibles cuando el tipo de pregunta las requiere.'),
('Respuesta','Lo que contestó el alumno en una pregunta de su intento.')
],[3.3,13.3])
p('Separar actividad, publicación e intento permite distinguir el examen preparado por el docente, su disponibilidad para una comisión y el trabajo realizado por cada alumno.').paragraph_format.space_before = Pt(7)
h('El recorrido de una evaluación')
for txt in [
'El docente prepara el contenido y lo guarda como borrador. Puede crear preguntas o traerlas desde su banco privado. Cada importación genera una copia independiente; los cambios posteriores no se sincronizan automáticamente con el original.',
'Cuando decide publicar, se envían el contenido final y la configuración a la API. La publicación completa es transaccional. Si el docente elige guardar todo en el banco, esa acción forma parte de la misma operación.',
'El alumno inicia un intento. El sistema selecciona una variante por grupo y registra esa selección en IntentoPregunta, de modo que el examen conserve las mismas preguntas al volver a consultarlo.',
'Al responder, la web guarda primero una copia local pendiente y luego intenta enviarla. Al recuperar la conexión, sincroniza las respuestas que faltan. Antes de entregar, debe confirmar que la cola pendiente quedó vacía.',
'La finalización calcula la parte automática de preguntas de opción múltiple y verdadero o falso. Las respuestas que requieren corrección manual quedan a cargo del docente; después, los resultados alimentan las consultas y los reportes.'
]: p(txt,'List Number')
p('Para evitar publicaciones duplicadas ante reintentos, se utiliza una clave de idempotencia: identifica la misma operación aunque se envíe otra vez. Las revisiones del banco y los borradores permiten detectar si otro cambio dejó desactualizada la versión que se está editando.')

doc.add_page_break()
h('Sesión y cuidado de las respuestas')
p('El usuario ingresa con correo y contraseña. NextAuth consulta a la API y mantiene la sesión en una cookie cifrada y HttpOnly. Esto impide que el código JavaScript de la página lea directamente las credenciales guardadas allí. La información de sesión expuesta a la interfaz contiene la identidad y el rol, sin los tokens de la API.')
p('La web controla la navegación según el rol, pero la autorización definitiva se realiza en la API. Por ejemplo, ocultar un botón de edición no alcanza: el servidor debe rechazar la operación si el usuario no tiene permiso.')
p('Durante un examen, las respuestas pendientes se conservan en una cola local, llamada outbox. Si se pierde la conexión, la interfaz mantiene el intento visible y diferencia lo guardado de lo pendiente. Si IndexedDB no está disponible, el respaldo queda en memoria y se advierte que no debe cerrarse ni recargarse la pestaña. Este mecanismo ayuda a conservar respuestas; no convierte toda la aplicación en un sistema sin conexión.')
h('La aplicación móvil')
p('Examio Native complementa el proyecto con React Native, Expo y TypeScript. Usa Expo Router para navegar y Zustand para mantener el usuario en memoria. Su objetivo es que el alumno consulte desde el celular próximas evaluaciones, fechas y calificaciones.')
p('En su estado actual tiene un inicio de sesión simulado y una pantalla de próximas actividades con datos estáticos. Todavía no está conectada con la API y la sesión no se conserva al cerrar o recargar por completo la aplicación. Los recordatorios, el historial de notas y las notificaciones figuran como funciones pendientes.')
h('Pruebas y estado actual')
p('El frontend cuenta con Vitest y React Testing Library para probar lógica y componentes, MSW para simular respuestas HTTP y Playwright para recorrer flujos en un navegador. Los recorridos automatizados usan una API de prueba local, sin modificar datos académicos reales.')
p('La web ya integra el banco privado, los borradores, la publicación y los reportes. Siguen pendientes mejoras en la gestión de etiquetas y en la presentación de conflictos de edición. En los reportes, los gráficos por alumno representan la página visible; los del reporte por evaluación usan agregaciones globales de la API.')
h2('Documentación de referencia')
p('Para ampliar: AGENTS.md y README.md del frontend; docs/API.md para la API; docs/DATABASE.md para las entidades; PLAN.md para el trabajo pendiente; y README.md de examio-native para el alcance móvil. Este documento describe la organización lógica del sistema, sin definir una infraestructura de producción concreta.')
doc.save(out)
print(out.resolve())
